#!/usr/bin/env python3
"""Writes the .xls and .ppt fixtures: an .xlsx/.pptx made with openpyxl / python-pptx is
saved as .xls/.ppt by LibreOffice, and that file converted back to .xlsx/.pptx by LibreOffice
(what the Document view falls back to), so tests can compare the native readers against it.

xls-workbook.xls    two sheets: text, numbers, booleans, a date, formulas, an error, gaps
xls-workbook.xlsx   xls-workbook.xls converted by LibreOffice
ppt-deck.ppt        three slides: title, bullets with a line break, a table, grouped shapes, notes
ppt-deck.pptx       ppt-deck.ppt converted by LibreOffice

Needs LibreOffice (soffice), openpyxl, python-pptx and cargo (the ignored test
strip_fixture_preview in src/convert/ppt.rs removes the preview image from the .ppt).
"""
import datetime
import os
import shutil
import subprocess
import tempfile

import openpyxl
from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))


def workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Parts"
    ws.append(["Part", "Count", "Price", "In stock", "Total"])
    ws.append(["Bolt M3", 12, 0.1, True, "=B2*C2"])
    ws.append(["Washer, heavy", 3, 1.5, False, "=B3*C3"])
    ws.append(["Nut", -7, 1 / 3, True, "=B4*C4"])
    ws.append([])
    ws.append(["Sum", "=SUM(B2:B4)", None, None, "=SUM(E2:E4)"])
    ws.append(["Label", '=A2&" x"&B2', "=1/0", "=B2>10", '=IF(B3>=3,"a>=b","")'])
    ws["A8"] = "Released"
    ws["B8"] = datetime.date(2026, 10, 19)
    ws["A9"] = 1e20
    ws["B9"] = 123456789
    ws["C9"] = 2.5e-5
    notes = wb.create_sheet("Notes ü")
    notes["A1"] = "Grüße aus Köln – €"
    notes["C3"] = "far away"
    notes["AA3"] = 42
    notes["A5"] = "line one\nline two"
    wb.save(path)


def deck(path):
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Quarterly Review"
    s1.placeholders[1].text = "Team Parallax"
    s1.notes_slide.notes_text_frame.text = "Welcome everyone.\nKeep it short."

    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Agenda"
    body = s2.placeholders[1].text_frame
    body.text = "Results"
    p = body.add_paragraph()
    p.text = "Plans"
    p.level = 1
    body.add_paragraph().text = "Risks\vand mitigations"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(5), Inches(4), Inches(1)).table
    for r, row in enumerate([["Metric", "Value"], ["Revenue", "1.2 M€"]]):
        for c, text in enumerate(row):
            table.cell(r, c).text = text

    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Architecture"
    group = s3.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1)).text_frame.text = "Core"
    group.shapes.add_textbox(Inches(4), Inches(2), Inches(2), Inches(1)).text_frame.text = "App"
    s3.notes_slide.notes_text_frame.text = "Mention the CLI."
    # Unused layouts would each become a master in the .ppt.
    used = [s.slide_layout for s in prs.slides]
    layouts = prs.slide_masters[0].slide_layouts
    for layout in list(layouts):
        if layout not in used:
            layouts.remove(layout)
    prs.save(path)


def convert(src, target, outdir):
    subprocess.run(["soffice", "--headless", "--convert-to", target, "--outdir", outdir, src], check=True,
                   stdout=subprocess.DEVNULL)


with tempfile.TemporaryDirectory() as tmp:
    for name, make, legacy, ooxml in [("xls-workbook", workbook, "xls", "xlsx"), ("ppt-deck", deck, "ppt", "pptx")]:
        source = os.path.join(tmp, "source", f"{name}.{ooxml}")
        os.makedirs(os.path.dirname(source), exist_ok=True)
        make(source)
        convert(source, legacy, tmp)
        shutil.copy(os.path.join(tmp, f"{name}.{legacy}"), os.path.join(HERE, f"{name}.{legacy}"))
        if legacy == "ppt":
            # Drops the preview image LibreOffice always embeds (half a megabyte).
            subprocess.run(["cargo", "test", "-q", "-p", "parallax-core", "--lib", "strip_fixture_preview", "--",
                            "--ignored"], check=True, cwd=HERE)
        back = os.path.join(tmp, "back")
        convert(os.path.join(HERE, f"{name}.{legacy}"), ooxml, back)
        shutil.copy(os.path.join(back, f"{name}.{ooxml}"), os.path.join(HERE, f"{name}.{ooxml}"))
